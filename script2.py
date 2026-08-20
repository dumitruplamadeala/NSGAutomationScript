import argparse
import shutil
import time
from pathlib import Path

import pythoncom
import xlwings as xw
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from win32com.client import Dispatch


def get_used_size(sheet):
    used = sheet.api.UsedRange

    last_row = used.Row + used.Rows.Count - 1
    last_col = used.Column + used.Columns.Count - 1

    return last_row, last_col


def clear_template_data(sheet, last_col):
    template_last_row, template_last_col = get_used_size(sheet)

    if template_last_row < 2:
        return

    last_col = max(last_col, template_last_col)

    sheet.api.Range(
        sheet.api.Cells(2, 1),
        sheet.api.Cells(template_last_row, last_col),
    ).ClearContents()


def copy_values_and_formulas(
    source_sheet,
    target_sheet,
    last_row,
    last_col,
):
    if last_row < 2:
        return

    source_range = source_sheet.api.Range(
        source_sheet.api.Cells(2, 1),
        source_sheet.api.Cells(last_row, last_col),
    )

    target_range = target_sheet.api.Range(
        target_sheet.api.Cells(2, 1),
        target_sheet.api.Cells(last_row, last_col),
    )

    target_range.Formula = source_range.Formula


def get_characters(cell, start, length):
    range_api = cell.api

    dispid = range_api._oleobj_.GetIDsOfNames("Characters")

    characters = range_api._oleobj_.Invoke(
        dispid,
        0,
        pythoncom.DISPATCH_PROPERTYGET,
        1,
        start,
        length,
    )

    return Dispatch(characters)


def apply_rich_text_run(target_cell, start, length, font):
    """
    Apply one rich-text run from openpyxl to the target Excel cell.
    """
    target_font = get_characters(
        target_cell,
        start,
        length,
    ).Font

    if font.b is not None:
        target_font.Bold = font.b

    if font.i is not None:
        target_font.Italic = font.i

    if font.u is not None:
        target_font.Underline = font.u

    if font.strike is not None:
        target_font.Strikethrough = font.strike

    if font.vertAlign == "subscript":
        target_font.Subscript = True

    if font.vertAlign == "superscript":
        target_font.Superscript = True


def copy_rich_text(
    rich_sheet,
    target_sheet,
    last_row,
    last_col,
):
    """
    Copy:
      - mixed rich-text runs
      - whole-cell bold formatting

    Everything else keeps the template formatting.
    """

    rich_cells = 0
    rich_runs = 0
    whole_bold_cells = 0

    for row in rich_sheet.iter_rows(
        min_row=2,
        max_row=last_row,
        min_col=1,
        max_col=last_col,
    ):
        for cell in row:

            target_cell = target_sheet.cells(
                cell.row,
                cell.column,
            )

            # -------------------------------------------------
            # Mixed rich text
            # -------------------------------------------------

            if isinstance(cell.value, CellRichText):

                position = 1

                for part in cell.value:

                    if isinstance(part, TextBlock):
                        text = part.text

                        if text:
                            apply_rich_text_run(
                                target_cell,
                                position,
                                len(text),
                                part.font,
                            )

                            rich_runs += 1

                    else:
                        text = str(part)

                    position += len(text)

                rich_cells += 1

            # -------------------------------------------------
            # Whole-cell bold
            # -------------------------------------------------

            elif (
                isinstance(cell.value, str)
                and cell.value
                and cell.font.bold is True
            ):
                target_cell.api.Font.Bold = True
                whole_bold_cells += 1

    print(
        f"  Rich text cells: {rich_cells}, "
        f"runs applied: {rich_runs}, "
        f"whole-cell bold: {whole_bold_cells}"
    )

def set_action_wrap_text(sheet, last_row, last_col):
    if last_row < 2:
        return

    headers = [
        sheet.cells(1, col).value
        for col in range(1, last_col + 1)
    ]

    try:
        action_col = next(
            col
            for col, header in enumerate(headers, start=1)
            if str(header).strip().casefold() == "action"
        )
    except StopIteration:
        print(
            f"Warning: Action column not found on '{sheet.name}'."
        )
        return

    sheet.api.Range(f"2:{last_row}").WrapText = False

    for row in range(2, last_row + 1):
        action = sheet.cells(row, action_col).value

        should_wrap = (
            str(action).strip().casefold()
            in {
                "create",
                "remove",
                "update",
            }
        )

        if should_wrap:
            sheet.api.Rows(row).WrapText = True


def merge_workbooks(
    source_file,
    template_file,
    output_file,
):
    source_file = Path(source_file).resolve()
    template_file = Path(template_file).resolve()
    output_file = Path(output_file).resolve()

    if output_file.exists():
        output_file.unlink()

    shutil.copy2(
        template_file,
        output_file,
    )

    # Open source once with openpyxl specifically for rich text.
    # We never save this workbook with openpyxl.
    rich_wb = load_workbook(
        source_file,
        rich_text=True,
        data_only=False,
    )

    app = None
    wb_source = None
    wb_output = None

    try:
        app = xw.App(
            visible=False
        )

        app.display_alerts = False
        app.screen_updating = False

        wb_source = app.books.open(
            str(source_file),
            read_only=True,
            update_links=False,
        )

        wb_output = app.books.open(
            str(output_file),
            update_links=False,
        )

        template_sheets = {
            sheet.name
            for sheet in wb_output.sheets
        }

        for source_sheet in wb_source.sheets:

            sheet_name = source_sheet.name

            if sheet_name not in template_sheets:
                print(
                    f"Skipping '{sheet_name}': "
                    f"not present in template."
                )
                continue

            print(
                f"Processing: {sheet_name}"
            )

            target_sheet = (
                wb_output.sheets[sheet_name]
            )

            rich_sheet = rich_wb[sheet_name]

            last_row, last_col = get_used_size(
                source_sheet
            )

            print(
                f"  Source range: "
                f"A1 to R{last_row}C{last_col}"
            )

            clear_template_data(
                target_sheet,
                last_col,
            )

            stage_start = time.perf_counter()

            copy_values_and_formulas(
                source_sheet,
                target_sheet,
                last_row,
                last_col,
            )

            print(
                f"  Values and formulas: "
                f"{time.perf_counter() - stage_start:.2f}s"
            )

            stage_start = time.perf_counter()

            copy_rich_text(
                rich_sheet,
                target_sheet,
                last_row,
                last_col,
            )

            print(
                f"  Rich text: "
                f"{time.perf_counter() - stage_start:.2f}s"
            )

            stage_start = time.perf_counter()

            set_action_wrap_text(
                target_sheet,
                last_row,
                last_col,
            )

            print(
                f"  Action wrap text: "
                f"{time.perf_counter() - stage_start:.2f}s"
            )

        wb_output.save()

    finally:
        rich_wb.close()

        if wb_source:
            try:
                wb_source.close()
            except Exception:
                pass

        if wb_output:
            try:
                wb_output.close()
            except Exception:
                pass

        if app:
            try:
                app.quit()
            except Exception:
                pass


def main():
    start_time = time.time()

    parser = argparse.ArgumentParser(
        description=(
            "Apply Excel template while preserving rich text."
        )
    )

    parser.add_argument(
        "input_file",
        help="Source Excel file",
    )

    parser.add_argument(
        "--template",
        default="assets/template.xlsx",
        help="Template file",
    )

    args = parser.parse_args()

    input_path = Path(
        args.input_file
    ).resolve()

    template_path = Path(
        args.template
    ).resolve()

    if not input_path.is_file():
        parser.error(
            f"Input file not found: {input_path}"
        )

    if not template_path.is_file():
        parser.error(
            f"Template file not found: {template_path}"
        )

    output_path = input_path.with_name(
        f"{input_path.stem}_templated"
        f"{template_path.suffix}"
    )

    merge_workbooks(
        input_path,
        template_path,
        output_path,
    )

    print(
        f"Created: {output_path.name}"
    )

    print(
        f"Finished in "
        f"{time.time() - start_time:.2f}s"
    )


if __name__ == "__main__":
    main()